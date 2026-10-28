"""Factual collection counts shared by terminal and Word presentation.

Unknown counts stay unknown. Source totals never become confidence scores or
claim-label percentages, and an unusable bundle cannot look like success.
"""

from __future__ import annotations

import textwrap
from dataclasses import dataclass
from typing import TYPE_CHECKING, Protocol

if TYPE_CHECKING:
    from docx.document import Document
    from docx.text.paragraph import Paragraph


class CoverageConsole(Protocol):
    @property
    def term_width(self) -> int:
        """Return the available output width in columns."""

    @property
    def INFO(self) -> str:
        """Return the accent escape sequence, or an empty string."""

    @property
    def WARNING(self) -> str:
        """Return the warning escape sequence, or an empty string."""

    @property
    def RESET(self) -> str:
        """Return the reset escape sequence, or an empty string."""

    def text(self, msg: str) -> None:
        """Write one result line on the console output stream."""


@dataclass(frozen=True)
class CollectionCoverage:
    """Recorded counts from one explicitly bound collection, never live guesses."""

    pages_captured: int | None = None
    first_party_pages: int | None = None
    blocked_pages: int | None = None
    redirect_loops: int | None = None
    hiring_postings: int | None = None
    usable: bool | None = None

    def __post_init__(self) -> None:
        for value in (
            self.pages_captured,
            self.first_party_pages,
            self.blocked_pages,
            self.redirect_loops,
            self.hiring_postings,
        ):
            if value is not None and (
                isinstance(value, bool) or not isinstance(value, int) or value < 0
            ):
                raise ValueError("Coverage counts must be non-negative integers or unknown")
        if (
            self.pages_captured is not None
            and self.first_party_pages is not None
            and self.first_party_pages > self.pages_captured
        ):
            raise ValueError("First-party pages cannot exceed all captured pages")

    @property
    def needs_attention(self) -> bool:
        return self.usable is False or self.pages_captured == 0

    def description(self) -> str:
        if self.usable is False:
            count = self.pages_captured if self.pages_captured is not None else "unknown"
            noun = "page" if self.pages_captured == 1 else "pages"
            pages = f"{count} {noun} captured (bundle not usable)"
        elif self.pages_captured == 0:
            pages = "0 pages captured"
        elif self.first_party_pages is not None:
            noun = "page" if self.first_party_pages == 1 else "pages"
            pages = f"{self.first_party_pages} first-party {noun}"
            if self.pages_captured is not None:
                other = self.pages_captured - self.first_party_pages
                if other:
                    noun = "page" if other == 1 else "pages"
                    pages += f", {other} other captured {noun}"
        elif self.pages_captured is not None:
            noun = "page" if self.pages_captured == 1 else "pages"
            pages = f"{self.pages_captured} {noun} captured"
        else:
            pages = "collection counts not recorded"
        if self.blocked_pages is not None:
            pages += f", {self.blocked_pages} blocked"
        if self.redirect_loops is not None:
            pages += f", {self.redirect_loops} redirect {'loop' if self.redirect_loops == 1 else 'loops'}"
        if self.hiring_postings is not None:
            noun = "posting" if self.hiring_postings == 1 else "postings"
            pages += f"; hiring: {self.hiring_postings} linked {noun}"
        return pages


def show_collection_coverage(summary: CollectionCoverage, console: CoverageConsole) -> None:
    """Print one logical signature with hanging continuation lines."""
    lines = textwrap.wrap(
        "coverage  " + summary.description(),
        width=max(20, console.term_width),
        subsequent_indent="          ",
        break_on_hyphens=False,
    )
    for index, line in enumerate(lines):
        if summary.needs_attention:
            if index == 0:
                line = f"{console.INFO}coverage{console.RESET}  {console.WARNING}{line[10:]}{console.RESET}"
            else:
                line = f"{console.WARNING}{line}{console.RESET}"
        else:
            if index == 0:
                line = line.replace("coverage", f"{console.INFO}coverage{console.RESET}", 1)
            if summary.blocked_pages:
                line = line.replace("blocked", f"{console.WARNING}blocked{console.RESET}")
            if summary.redirect_loops:
                for word in ("redirect", "loop" if summary.redirect_loops == 1 else "loops"):
                    line = line.replace(word, f"{console.WARNING}{word}{console.RESET}")
        console.text(line)


def add_docx_collection_note(document: Document, summary: CollectionCoverage) -> Paragraph:
    """Append the recorded counts as a discreet closing methodology note.

    Client-facing reports open with the title and analysis; collection counts
    belong after the content, never under the title. The wording and meaning
    match the terminal signature. A bundle that needs attention keeps the
    warning color so moving the note never hides unreliable collection.
    """
    from docx.shared import Pt, RGBColor

    from primr.output.style_engine import StyleEngine

    paragraph = document.add_paragraph()
    label = paragraph.add_run("Collection note: ")
    label.font.size = Pt(9)
    label.font.bold = True
    label.font.color.rgb = StyleEngine.MUTED_COLOR
    tokens = (summary.description() + ".").split(" ")
    for index, token in enumerate(tokens):
        run = paragraph.add_run(token if index == len(tokens) - 1 else token + " ")
        run.font.size = Pt(9)
        warning = (
            summary.needs_attention
            or (summary.blocked_pages and token.startswith("blocked"))
            or (summary.redirect_loops and token.rstrip(";,.") in {"redirect", "loop", "loops"})
        )
        run.font.color.rgb = RGBColor(0x85, 0x55, 0x00) if warning else StyleEngine.MUTED_COLOR
    return paragraph
