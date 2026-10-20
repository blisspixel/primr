"""Evidence labels and actual text/background contrast in generated Word files."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

from primr.output.document_builder import DocumentBuilder
from primr.output.polish_elements import DataConfidenceIndicator, StrategicRecommendationFormatter
from primr.output.style_engine import DualCodingEnhancer, StyleEngine


def _contrast(first, second):
    def luminance(color):
        channels = [value / 255 for value in color]
        linear = [
            value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4
            for value in channels
        ]
        return sum(
            value * weight for value, weight in zip(linear, (0.2126, 0.7152, 0.0722), strict=True)
        )

    bright, dark = sorted((luminance(first), luminance(second)), reverse=True)
    return (bright + 0.05) / (dark + 0.05)


@pytest.mark.parametrize(
    "foreground,background",
    [
        (StyleEngine.PRIMARY_COLOR, (255, 255, 255)),
        (StyleEngine.SECONDARY_COLOR, (255, 255, 255)),
        (StyleEngine.ACCENT_COLOR, (255, 255, 255)),
        (StyleEngine.TEXT_COLOR, StyleEngine.TABLE_ALT_ROW),
        (StyleEngine.TEXT_COLOR, StyleEngine.HIGHLIGHT_BG),
        (StyleEngine.MUTED_COLOR, (255, 255, 255)),
        (StyleEngine.MUTED_COLOR, StyleEngine.HIGHLIGHT_BG),
        ((255, 255, 255), StyleEngine.TABLE_HEADER_BG),
    ],
)
def test_report_text_background_pairs_meet_normal_text_contrast(foreground, background):
    assert _contrast(foreground, background) >= 4.5


def test_saved_docx_preserves_each_risks_evidence_without_guessing():
    risks = [
        "(Reported, S004) Customer concentration poses a material risk.",
        "(Hypothesis) Distribution risk may limit expansion.",
        "Supply chain disruption remains a material risk.",
    ]
    builder = DocumentBuilder("ExampleCo", {"board_of_directors_concerns": "\n\n".join(risks)})
    document = builder.build()
    target = BytesIO()
    document.save(target)
    reloaded = Document(target)
    paragraphs = [p.text for p in reloaded.paragraphs]
    watch = paragraphs.index("Watch Outs")
    assert paragraphs[watch + 1 : watch + 4] == risks
    text = "\n".join(paragraphs)
    for mark in ("●●", "●○", "Data Confidence:"):
        assert mark not in text


@pytest.mark.parametrize("score", ["high", "medium", "low", "unknown"])
def test_legacy_generic_confidence_scores_do_not_render_claim_marks(score):
    indicator = DataConfidenceIndicator()
    assert indicator.get_indicator(score) == ""
    assert indicator.format_with_confidence("Claim", score) == "Claim"
    assert DualCodingEnhancer().add_confidence_indicator("Claim", score) == "Claim"


@pytest.mark.parametrize("label", ["Confirmed", "Reported", "Estimated", "Hypothesis"])
def test_compatibility_formatter_accepts_only_explicit_evidence_words(label):
    indicator = DataConfidenceIndicator()
    text = f"Claim ({label}, S004)"
    assert indicator.format_with_confidence(text, label) == text
    assert indicator.format_with_confidence("Claim", label) == f"Claim ({label})"


def test_recommendation_priority_uses_words_in_navy():
    document = Document()
    formatter = StrategicRecommendationFormatter(document)
    for index, priority in enumerate(("high", "medium", "low"), 1):
        formatter.format_recommendation(index, "Review the evidence.", priority)
    for paragraph, word in zip(document.paragraphs, ("[HIGH]", "[MED]", "[LOW]"), strict=True):
        assert word in paragraph.text
        assert paragraph.runs[1].font.color.rgb == StyleEngine.PRIMARY_COLOR


def test_saved_document_has_accessible_explicit_text_colors():
    builder = DocumentBuilder(
        "ExampleCo",
        {
            "company_overview": "ExampleCo serves enterprise customers (Reported, S001).",
            "strategic_recommendations": "Critical supply risk should be reviewed (Hypothesis).",
            "financial_overview": "| Metric | Value |\n|---|---|\n| Revenue | Not disclosed |",
        },
        citations=[{"title": "Source", "url": "https://example.co"}],
    )
    document = builder.build()
    StrategicRecommendationFormatter(document).format_recommendation(1, "Review.", "high")
    target = BytesIO()
    document.save(target)
    reloaded = Document(target)
    # Check explicit run colors in the actual body, tables, headers and footers.
    # Style inheritance is covered by the palette-pair checks above.
    elements = [
        reloaded.element,
        *(section.header._element for section in reloaded.sections),
        *(section.footer._element for section in reloaded.sections),
    ]
    for element in elements:
        for run in element.iter(qn("w:r")):
            for color in run.iter(qn("w:color")):
                value = color.get(qn("w:val"))
                if not value or value == "auto":
                    continue
                rgb = tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))
                background = (255, 255, 255)
                for ancestor in run.iterancestors():
                    if ancestor.tag == qn("w:tc"):
                        shading = ancestor.find("w:tcPr/w:shd", namespaces=ancestor.nsmap)
                        if shading is not None:
                            fill = shading.get(qn("w:fill"))
                            if fill and fill != "auto":
                                background = tuple(
                                    int(fill[index : index + 2], 16) for index in (0, 2, 4)
                                )
                        break
                assert _contrast(rgb, background) >= 4.5, (value, background)


def test_standalone_markdown_renderer_uses_the_accessible_report_palette(tmp_path):
    from primr.output.markdown_converter import markdown_to_docx

    target = tmp_path / "ExampleCo.docx"
    markdown_to_docx(
        "# ExampleCo\n\n## Evidence\n\n### Details\n\n#### Next step\n\nRead the source.", target
    )
    document = Document(target)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        rgb = document.styles[name].font.color.rgb
        assert rgb is not None, name
        assert _contrast(tuple(rgb), (255, 255, 255)) >= 4.5, name
