"""One factual signature, with no invented counts or confidence claims."""

from __future__ import annotations

import re
from io import BytesIO

import pytest
from docx import Document

from primr.output.collection_coverage import (
    CollectionCoverage,
    add_docx_coverage,
    show_collection_coverage,
)
from primr.utils.console import Console, TerminalCapabilities


def _console():
    return Console(
        capabilities=TerminalCapabilities.for_testing(
            supports_color=True,
            supports_unicode=False,
            supports_cursor=True,
            width=60,
            is_interactive=True,
        )
    )


def test_counts_wrap_once_with_hanging_indent(capsys):
    summary = CollectionCoverage(14, 14, 2, 1, 3, True)
    show_collection_coverage(summary, _console())
    text = capsys.readouterr().out
    plain = re.sub(r"\x1b\[[0-9;]*m", "", text)
    assert plain.count("coverage") == 1
    assert all(len(line) <= 60 for line in plain.splitlines())
    assert all(line.startswith("          ") for line in plain.splitlines()[1:])
    assert "14 first-party pages" in plain
    assert "2 blocked" in plain
    assert "\x1b[33mblocked" in text
    assert "3 linked postings" in plain
    assert "%" not in text


@pytest.mark.parametrize("pages", [0, 2])
def test_unusable_bundle_is_yellow_even_with_captured_pages(pages, capsys):
    show_collection_coverage(CollectionCoverage(pages_captured=pages, usable=False), _console())
    text = capsys.readouterr().out
    assert "bundle not usable" in text
    assert f"\x1b[33m{pages} pages captured" in text


def test_zero_blocked_counts_never_use_warning_color(capsys):
    show_collection_coverage(CollectionCoverage(2, 2, 0, 0, 0, True), _console())
    text = capsys.readouterr().out
    assert "\x1b[33m" not in text


def test_unknown_counts_do_not_become_zero():
    text = CollectionCoverage().description()
    assert text == "collection counts not recorded"
    assert "0" not in text


@pytest.mark.parametrize("value", [-1, True, "1", float("nan")])
def test_invalid_counts_are_rejected(value):
    with pytest.raises(ValueError):
        CollectionCoverage(pages_captured=value)


def test_public_or_archived_sources_are_not_counted_as_first_party():
    text = CollectionCoverage(5, 2).description()
    assert text == "2 first-party pages, 3 other captured pages"
    with pytest.raises(ValueError):
        CollectionCoverage(2, 3)


def test_word_signature_is_the_same_once_under_title():
    summary = CollectionCoverage(14, 14, 2, 1, 3, True)
    document = Document()
    document.add_heading("ExampleCo", level=0)
    add_docx_coverage(document, summary)
    buffer = BytesIO()
    document.save(buffer)
    loaded = Document(buffer)
    assert loaded.paragraphs[1].text == "coverage  " + summary.description()
    assert sum("coverage" in paragraph.text for paragraph in loaded.paragraphs) == 1
