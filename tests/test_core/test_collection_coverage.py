"""Collection signatures use their owned inventory, including failure controls."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from primr.core.cli_research_result import finalize_research_command
from primr.core.collection_coverage import (
    coverage_from_index,
    load_diagnostics_coverage,
    load_run_coverage,
    record_run_collection,
)
from primr.output.collection_coverage import CollectionCoverage
from primr.output.document_builder import DocumentBuilder
from primr.output.markdown_converter import markdown_to_docx


def _record(root: Path):
    hiring = root / "_hiring"
    hiring.mkdir()
    (hiring / "postings_index.json").write_text(
        json.dumps(
            [
                {"url": "https://jobs.acme.example/1", "identity": "verified"},
                {"url": "https://jobs.other.example/2", "identity": "candidate"},
            ]
        )
    )
    raw = root / "_raw_scrapes"
    raw.mkdir()
    (raw / "fb_old.txt").write_text("URL: https://acme.example/old\nSource: wayback\n")
    record_run_collection(
        str(root),
        "https://acme.example",
        {
            "https://acme.example/about": "Current first-party capture.",
            "https://acme.example/old": "Historical archived capture.",
            "https://other.example/review": "Independent public capture.",
            "https://acme.example/empty": " ",
        },
    )


def test_owned_index_and_manifest_match_without_claiming_unknown_trace_counts(tmp_path):
    _record(tmp_path)
    summary = load_run_coverage(tmp_path)
    assert summary == CollectionCoverage(pages_captured=3, first_party_pages=1, hiring_postings=1)
    state = json.loads((tmp_path / "_run_state.json").read_text())
    for key in ("pages_captured", "first_party_pages", "hiring_postings"):
        assert state["collection_coverage"][key] == getattr(summary, key)
    assert summary.blocked_pages is None
    assert summary.redirect_loops is None
    assert load_diagnostics_coverage(tmp_path / "_diagnostics") == summary
    assert load_diagnostics_coverage(tmp_path / "other") is None


@pytest.mark.parametrize("mutation", ["index", "state", "missing", "symlink"])
def test_unbound_or_tampered_inventory_never_becomes_success(tmp_path, mutation):
    _record(tmp_path)
    index = tmp_path / "source_index.json"
    if mutation == "index":
        data = json.loads(index.read_text())
        data["sources"].pop()
        index.write_text(json.dumps(data))
    elif mutation == "state":
        path = tmp_path / "_run_state.json"
        state = json.loads(path.read_text())
        state["collection_coverage"]["pages_captured"] = 99
        path.write_text(json.dumps(state))
    elif mutation == "missing":
        index.unlink()
    else:
        target = tmp_path / "elsewhere.json"
        index.rename(target)
        try:
            index.symlink_to(target)
        except OSError:
            pytest.skip("Symlinks unavailable on this platform")
    assert load_run_coverage(tmp_path) == CollectionCoverage()


def test_unknown_schema_duplicate_identity_and_bad_rows_are_rejected():
    with pytest.raises(ValueError):
        coverage_from_index({})
    row = {"source_id": "S001", "source_type": "first_party", "captured_characters": 10}
    for rows in (
        [row, row],
        [dict(row, captured_characters=True)],
        [dict(row, source_type="invented")],
        [None],
    ):
        with pytest.raises(ValueError):
            coverage_from_index(
                {"schema": "primr.evidence-source-index", "version": "1.0", "sources": rows}
            )


@pytest.mark.parametrize("title", [None, "ExampleCo"])
def test_saved_markdown_opens_with_title_and_summary_and_ends_with_the_note(tmp_path, title):
    path = tmp_path / "report.docx"
    summary = CollectionCoverage(3, 1, hiring_postings=1)
    body = "## Executive Summary\n\nEvidence remains.\n"
    markdown = f"# {title}\n\n{body}" if title else body
    markdown_to_docx(markdown, path, title=title, coverage=summary)
    paragraphs = [p.text for p in Document(path).paragraphs if p.text.strip()]
    opening = ["Executive Summary", "Evidence remains."]
    assert paragraphs[:-1] == ([title, *opening] if title else opening)
    assert paragraphs[-1] == f"Collection note: {summary.description()}."
    assert not any(p.startswith("coverage") for p in paragraphs)


def test_premium_cover_has_no_counts_and_the_note_closes_the_document(tmp_path):
    summary = CollectionCoverage(3, 1, hiring_postings=1)
    builder = DocumentBuilder("ExampleCo", {}, coverage=summary)
    path = tmp_path / "premium.docx"
    builder.build().save(path)
    paragraphs = [p.text for p in Document(path).paragraphs if p.text.strip()]
    assert paragraphs[0] == "ExampleCo"
    assert paragraphs[-1] == f"Collection note: {summary.description()}."
    assert sum(summary.description() in p for p in paragraphs) == 1


@pytest.mark.parametrize("json_output", [False, True])
def test_completion_signature_precedes_next_and_never_pollutes_json(tmp_path, capsys, json_output):
    _record(tmp_path)
    report = tmp_path / "report.md"
    report.write_text("# ExampleCo")
    finalize_research_command(
        result_path=str(report),
        run_context={"working_folder": str(tmp_path)},
        company_name="ExampleCo",
        website="https://acme.example",
        mode="full",
        json_output=json_output,
        open_after=False,
        open_result=lambda path: None,
    )
    output = capsys.readouterr().out
    if json_output:
        json.loads(output)
        assert "coverage  " not in output
        assert "Next:" not in output
    else:
        assert output.count("coverage  ") == 1
        assert output.index("coverage  ") < output.index("Next:")
        assert "1 first-party page" in output


def test_render_refuses_a_missing_explicit_bundle_before_writing_artifacts(tmp_path):
    from primr.core.cli_render import run_render

    report = tmp_path / "report.md"
    report.write_text("# ExampleCo\n\nA sourced report.")
    assert run_render(["render", str(report), "--evidence-bundle", str(tmp_path / "missing")]) == 1
    assert not report.with_suffix(".docx").exists()
    assert not report.with_suffix(".txt").exists()


def test_text_word_path_opens_with_title_and_ends_with_the_note(tmp_path):
    from primr.output.output_utils import save_report_as_docx

    summary = CollectionCoverage(4, 2, blocked_pages=1)
    txt = tmp_path / "report.txt"
    txt.write_text(
        "ExampleCo Strategic Overview\nOctober 27, 2026\n\n## Executive Summary\n\nBody.\n",
        encoding="utf-8",
    )
    path = save_report_as_docx(txt, "ExampleCo", output_dir=tmp_path, coverage=summary)
    paragraphs = [p.text for p in Document(path).paragraphs if p.text.strip()]
    assert paragraphs[:2] == ["ExampleCo Strategic Overview", "October 27, 2026"]
    assert paragraphs[-1] == f"Collection note: {summary.description()}."
    assert sum(summary.description() in p for p in paragraphs) == 1
